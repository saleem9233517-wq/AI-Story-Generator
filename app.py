from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    send_file
)

from openai import OpenAI
from dotenv import load_dotenv

import os
import io
import base64
import json
import re
import pyttsx3

from database import (
    create_database,
    save_story,
    get_stories,
    get_story,
    update_story,
    update_cover,
    delete_story
)

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

from docx import Document


# =====================================================
# LOAD ENVIRONMENT
# =====================================================

load_dotenv()


# =====================================================
# FLASK APP
# =====================================================

app = Flask(__name__)


# =====================================================
# OPENAI
# =====================================================

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)


# =====================================================
# CREATE DATABASE
# =====================================================

create_database()


# =====================================================
# CREATE COVER FOLDER
# =====================================================

COVER_FOLDER = os.path.join(
    "static",
    "covers"
)

os.makedirs(
    COVER_FOLDER,
    exist_ok=True
)


# =====================================================
# CREATE AUDIO FOLDER
# =====================================================

AUDIO_FOLDER = os.path.join(
    "static",
    "audio"
)

os.makedirs(
    AUDIO_FOLDER,
    exist_ok=True
)


# =====================================================
# CREATE SCENE FOLDER
# =====================================================

SCENE_FOLDER = os.path.join(
    "static",
    "scenes"
)

os.makedirs(
    SCENE_FOLDER,
    exist_ok=True
)


# =====================================================
# GENERATE TITLE
# =====================================================

def generate_title(prompt, genre):

    title_prompt = f"""
    Create one creative and attractive title for this story.

    Story Idea:
    {prompt}

    Genre:
    {genre}

    Requirements:
    - Give only the title.
    - Do not add quotation marks.
    - Keep it short and memorable.
    - Make it suitable for the story.
    """

    response = client.responses.create(
        model="gpt-5.6-luna",
        input=title_prompt
    )

    title = response.output_text.strip()

    return title


# =====================================================
# GENERATE AI STORY
# =====================================================

def generate_story(
    prompt,
    genre,
    length,
    tone,
    title,
    character_name="",
    number_of_characters="1",
    setting="",
    writing_style="Creative",
    language="English",
    custom_instructions=""
):

    ai_prompt = f"""
    Write a creative and engaging story.

    Story Idea:
    {prompt}

    Genre:
    {genre}

    Length:
    {length}

    Tone:
    {tone}

    Story Title:
    {title}

    Character Name:
    {character_name if character_name else "Create a suitable main character"}

    Number of Main Characters:
    {number_of_characters}

    Story Setting:
    {setting if setting else "Choose a suitable setting based on the story idea"}

    Writing Style:
    {writing_style}

    Story Language:
    {language}

    Custom Instructions:
    {custom_instructions if custom_instructions else "No additional instructions"}

    Requirements:
    - Write the complete story.
    - Use the given title as inspiration.
    - Make the story interesting and original.
    - Follow the requested genre.
    - Follow the requested tone.
    - Match the requested length.
    - Use the requested character name if provided.
    - Use approximately the requested number of main characters.
    - Use the requested setting.
    - Follow the requested writing style.
    - Write the story in the requested language.
    - Follow the custom instructions if provided.
    - Do not generate another title.
    - Write only the story.
    """

    response = client.responses.create(
        model="gpt-5.6-luna",
        input=ai_prompt
    )

    return response.output_text


# =====================================================
# GENERATE AI STORY COVER
# =====================================================

def generate_cover(
    title,
    prompt,
    genre,
    tone,
    story_id
):

    cover_prompt = f"""
    Create a beautiful cinematic digital story cover
    based specifically on the following story.

    Story Title:
    {title}

    Story Idea:
    {prompt}

    Genre:
    {genre}

    Tone:
    {tone}

    Instructions:

    - The image MUST be strongly related to the story idea.
    - Show the main visual concept of the story.
    - Match the selected genre.
    - Match the selected tone.
    - Use cinematic composition.
    - Make it look like a professional story book cover
      illustration.
    - High quality and detailed.
    - Attractive lighting.
    - Create a visually interesting scene.
    - Do NOT use a generic placeholder.
    - Do NOT include text.
    - Do NOT include letters.
    - Do NOT include the story title.
    - Do NOT include typography.
    """

    try:

        response = client.responses.create(
            model="gpt-5.6-luna",
            input=cover_prompt,
            tools=[
                {
                    "type": "image_generation"
                }
            ]
        )

        image_data = None

        for output in response.output:

            if output.type == "image_generation_call":

                image_data = output.result

                break

        if not image_data:

            print("No image returned from AI.")

            return None

        image_bytes = base64.b64decode(
            image_data
        )

        filename = (
            f"story_{story_id}.png"
        )

        file_path = os.path.join(
            COVER_FOLDER,
            filename
        )

        with open(
            file_path,
            "wb"
        ) as image_file:

            image_file.write(
                image_bytes
            )

        if not os.path.exists(file_path):

            print(
                "ERROR: Cover file was not created."
            )

            return None

        file_size = os.path.getsize(
            file_path
        )

        if file_size == 0:

            print(
                "ERROR: Cover file is empty."
            )

            return None

        cover_path = (
            f"covers/{filename}"
        )

        print(
            "AI COVER CREATED:",
            cover_path
        )

        return cover_path

    except Exception as e:

        print(
            "AI COVER ERROR:",
            type(e).__name__,
            str(e)
        )

        return None


# =====================================================
# SPLIT STORY INTO AUDIO CHUNKS
# =====================================================

def split_story_for_voice(
    text,
    max_chars=3500
):

    words = text.split()

    chunks = []

    current_chunk = ""

    for word in words:

        test_chunk = (
            current_chunk + " " + word
        ).strip()

        if len(test_chunk) <= max_chars:

            current_chunk = test_chunk

        else:

            if current_chunk:

                chunks.append(
                    current_chunk
                )

            current_chunk = word

    if current_chunk:

        chunks.append(
            current_chunk
        )

    return chunks


# =====================================================
# GENERATE STORY VOICE
# =====================================================

def create_story_voice(
    story_id,
    story_text
):

    print("\n================================")
    print("VOICE GENERATION")
    print("================================")

    print(
        "Story ID:",
        story_id
    )

    print(
        "Story characters:",
        len(story_text)
    )

    if not story_text:

        raise ValueError(
            "Story text is empty."
        )

    os.makedirs(
        AUDIO_FOLDER,
        exist_ok=True
    )

    final_filename = (
        f"story_{story_id}.mp3"
    )

    final_path = os.path.join(
        AUDIO_FOLDER,
        final_filename
    )

    if os.path.exists(final_path):

        try:

            os.remove(final_path)

            print(
                "Old audio removed."
            )

        except Exception as e:

            print(
                "Could not remove old audio:",
                e
            )

    chunks = split_story_for_voice(
        story_text,
        3500
    )

    print(
        "Total voice chunks:",
        len(chunks)
    )

    try:

        from pydub import AudioSegment

    except ImportError:

        raise Exception(
            "pydub is not installed. Run: pip install pydub"
        )

    combined_audio = None

    for index, chunk in enumerate(
        chunks,
        start=1
    ):

        print(
            f"\nGenerating voice chunk {index}/{len(chunks)}..."
        )

        print(
            "Chunk characters:",
            len(chunk)
        )

        chunk_filename = (
            f"story_{story_id}_part_{index}.mp3"
        )

        chunk_path = os.path.join(
            AUDIO_FOLDER,
            chunk_filename
        )

        try:

            response = client.audio.speech.create(
                model="gpt-4o-mini-tts",
                voice="alloy",
                input=chunk
            )

            print(
                "OpenAI response received."
            )

            response.stream_to_file(
                chunk_path
            )

            print(
                "Chunk saved:",
                chunk_path
            )

            if not os.path.exists(
                chunk_path
            ):

                raise Exception(
                    "Audio chunk was not created."
                )

            file_size = os.path.getsize(
                chunk_path
            )

            print(
                "Chunk size:",
                file_size,
                "bytes"
            )

            if file_size == 0:

                raise Exception(
                    "Audio chunk is empty."
                )

            audio_segment = AudioSegment.from_mp3(
                chunk_path
            )

            if combined_audio is None:

                combined_audio = audio_segment

            else:

                combined_audio += audio_segment

        except Exception as e:

            print(
                "\nVOICE CHUNK ERROR:"
            )

            print(
                type(e).__name__
            )

            print(
                str(e)
            )

            raise

    if combined_audio is None:

        raise Exception(
            "No audio was generated."
        )

    print(
        "\nCreating final MP3..."
    )

    combined_audio.export(
        final_path,
        format="mp3"
    )

    for index in range(
        1,
        len(chunks) + 1
    ):

        chunk_filename = (
            f"story_{story_id}_part_{index}.mp3"
        )

        chunk_path = os.path.join(
            AUDIO_FOLDER,
            chunk_filename
        )

        if os.path.exists(
            chunk_path
        ):

            try:

                os.remove(
                    chunk_path
                )

            except Exception as e:

                print(
                    "Temporary audio delete error:",
                    e
                )

    if not os.path.exists(
        final_path
    ):

        raise Exception(
            "Final MP3 was not created."
        )

    final_size = os.path.getsize(
        final_path
    )

    print(
        "\n================================"
    )

    print(
        "VOICE GENERATED SUCCESSFULLY"
    )

    print(
        "File:",
        os.path.abspath(final_path)
    )

    print(
        "Size:",
        final_size,
        "bytes"
    )

    print(
        "================================\n"
    )

    return final_path


# =====================================================
# AI SCENE GENERATOR
# =====================================================

def generate_story_scenes(
    story_id,
    title,
    story_text,
    genre,
    tone
):

    print("\n================================")
    print("AI SCENE GENERATION STARTED")
    print("================================")

    print(
        "Story ID:",
        story_id
    )

    if not story_text:

        raise ValueError(
            "Story text is empty."
        )

    scene_prompt = f"""
    You are an expert film director and screenplay
    scene planner.

    Convert the following complete story into a
    cinematic scene plan for an AI-generated video.

    Story Title:
    {title}

    Genre:
    {genre}

    Tone:
    {tone}

    Complete Story:
    {story_text}

    IMPORTANT REQUIREMENTS:

    1. Divide the story into meaningful cinematic scenes.
    2. Do NOT remove important story events.
    3. Keep the story sequence logical.
    4. Each scene should represent one visual moment.
    5. Create between 5 and 12 scenes depending on story length.
    6. Each scene should be suitable for generating a video shot.
    7. Keep characters visually consistent between scenes.
    8. Clearly describe the location.
    9. Clearly describe character actions.
    10. Describe the camera shot.
    11. Describe lighting and atmosphere.
    12. Include approximate scene duration.
    13. Include narration text for the scene.
    14. Include dialogue only when dialogue naturally exists.
    15. Do not invent unrelated events.

    Return ONLY valid JSON.

    Use exactly this structure:

    {{
        "story_id": {story_id},
        "title": "{title}",
        "scenes": [
            {{
                "scene_number": 1,
                "scene_title": "Scene title",
                "location": "Location",
                "characters": [
                    "Character 1"
                ],
                "action": "What happens visually",
                "camera": "Camera shot and movement",
                "lighting": "Lighting and atmosphere",
                "duration": 8,
                "narration": "Narration for this scene",
                "dialogue": "Dialogue or empty string"
            }}
        ]
    }}

    Do not add markdown.
    Do not add ```json.
    Do not add explanations outside JSON.
    """

    response = client.responses.create(
        model="gpt-5.6-luna",
        input=scene_prompt
    )

    raw_output = response.output_text.strip()

    print(
        "\nAI SCENE RESPONSE RECEIVED"
    )

    clean_output = raw_output

    if clean_output.startswith(
        "```json"
    ):

        clean_output = clean_output[
            7:
        ]

        if clean_output.endswith(
            "```"
        ):

            clean_output = clean_output[
                :-3
            ]

    elif clean_output.startswith(
        "```"
    ):

        clean_output = clean_output[
            3:
        ]

        if clean_output.endswith(
            "```"
        ):

            clean_output = clean_output[
                :-3
            ]

    clean_output = clean_output.strip()

    if not clean_output.startswith(
        "{"
    ):

        match = re.search(
            r"\{.*\}",
            clean_output,
            re.DOTALL
        )

        if match:

            clean_output = match.group(0)

    try:

        scenes_data = json.loads(
            clean_output
        )

    except json.JSONDecodeError as e:

        print(
            "\nSCENE JSON ERROR:"
        )

        print(
            str(e)
        )

        print(
            "\nAI OUTPUT:"
        )

        print(
            raw_output
        )

        raise Exception(
            "AI returned invalid scene JSON."
        )

    if "scenes" not in scenes_data:

        raise Exception(
            "Scene data does not contain scenes."
        )

    if not isinstance(
        scenes_data["scenes"],
        list
    ):

        raise Exception(
            "Scenes data is not a list."
        )

    if len(
        scenes_data["scenes"]
    ) == 0:

        raise Exception(
            "No scenes were generated."
        )

    scene_filename = (
        f"story_{story_id}.json"
    )

    scene_path = os.path.join(
        SCENE_FOLDER,
        scene_filename
    )

    with open(
        scene_path,
        "w",
        encoding="utf-8"
    ) as scene_file:

        json.dump(
            scenes_data,
            scene_file,
            indent=4,
            ensure_ascii=False
        )

    print(
        "\n================================"
    )

    print(
        "AI SCENES GENERATED SUCCESSFULLY"
    )

    print(
        "Total scenes:",
        len(scenes_data["scenes"])
    )

    print(
        "Scene file:",
        os.path.abspath(scene_path)
    )

    print(
        "================================\n"
    )

    return scenes_data


# =====================================================
# GET SAVED SCENES
# =====================================================

def get_story_scenes(story_id):

    scene_filename = (
        f"story_{story_id}.json"
    )

    scene_path = os.path.join(
        SCENE_FOLDER,
        scene_filename
    )

    if not os.path.exists(
        scene_path
    ):

        return None

    try:

        with open(
            scene_path,
            "r",
            encoding="utf-8"
        ) as scene_file:

            scenes_data = json.load(
                scene_file
            )

        return scenes_data

    except Exception as e:

        print(
            "Scene file read error:",
            e
        )

        return None


# =====================================================
# GENERATE SCENE IMAGE
# =====================================================

def generate_scene_image(
    story_id,
    scene_number,
    scene
):

    print("\n==============================")
    print("SCENE IMAGE GENERATION")
    print("==============================")

    print(
        "Story ID:",
        story_id
    )

    print(
        "Scene:",
        scene_number
    )

    try:

        # =================================================
        # CREATE SCENE FOLDER
        # =================================================

        os.makedirs(
            SCENE_FOLDER,
            exist_ok=True
        )

        print(
            "Scene folder:",
            os.path.abspath(
                SCENE_FOLDER
            )
        )

        # =================================================
        # GET CHARACTERS
        # =================================================

        characters = scene.get(
            "characters",
            []
        )

        if isinstance(
            characters,
            list
        ):

            characters_text = ", ".join(
                str(character)
                for character in characters
            )

        else:

            characters_text = str(
                characters
            )

        # =================================================
        # IMAGE PROMPT
        # =================================================

        scene_prompt = f"""
        Create a cinematic fantasy movie scene.

        Scene title:
        {scene.get("scene_title", "")}

        Location:
        {scene.get("location", "")}

        Characters:
        {characters_text}

        Action:
        {scene.get("action", "")}

        Camera:
        {scene.get("camera", "")}

        Lighting and atmosphere:
        {scene.get("lighting", "")}

        Visual requirements:

        - cinematic fantasy film look
        - highly detailed environment
        - realistic characters
        - dramatic composition
        - consistent visual world
        - beautiful cinematic lighting
        - widescreen movie frame
        - professional film still
        - visually rich
        - no text
        - no subtitles
        - no watermark
        - no letters
        - no typography

        Create only the visual scene.

        Do not include narration or dialogue
        as text in the image.
        """

        print(
            "\nCalling image generation API..."
        )

        # =================================================
        # CALL IMAGE GENERATION
        # =================================================

        response = client.responses.create(
            model="gpt-5.6-luna",
            input=scene_prompt,
            tools=[
                {
                    "type": "image_generation"
                }
            ]
        )

        print(
            "Image API response received."
        )

        # =================================================
        # GET IMAGE DATA
        # =================================================

        image_data = None

        for output in response.output:

            print(
                "Output type:",
                getattr(
                    output,
                    "type",
                    "unknown"
                )
            )

            if output.type == "image_generation_call":

                image_data = output.result

                break

        # =================================================
        # CHECK IMAGE DATA
        # =================================================

        if not image_data:

            print(
                "ERROR: No image data returned by AI."
            )

            return None

        print(
            "Image data received."
        )

        # =================================================
        # DECODE IMAGE
        # =================================================

        try:

            image_bytes = base64.b64decode(
                image_data
            )

        except Exception as decode_error:

            print(
                "BASE64 DECODE ERROR:",
                type(decode_error).__name__,
                str(decode_error)
            )

            return None

        # =================================================
        # CHECK DECODED DATA
        # =================================================

        if not image_bytes:

            print(
                "ERROR: Decoded image data is empty."
            )

            return None

        print(
            "Decoded image bytes:",
            len(image_bytes)
        )

        # =================================================
        # IMAGE FILE NAME
        # =================================================

        image_filename = (
            f"story_{story_id}_scene_{scene_number}.png"
        )

        # =================================================
        # IMAGE PATH
        # =================================================

        image_path = os.path.join(
            SCENE_FOLDER,
            image_filename
        )

        print(
            "Image path:",
            os.path.abspath(
                image_path
            )
        )

        # =================================================
        # SAVE IMAGE
        # =================================================

        with open(
            image_path,
            "wb"
        ) as image_file:

            image_file.write(
                image_bytes
            )

        print(
            "Image write operation completed."
        )

        # =================================================
        # CHECK IMAGE EXISTS
        # =================================================

        if not os.path.exists(
            image_path
        ):

            print(
                "ERROR: Image file was NOT created."
            )

            return None

        # =================================================
        # CHECK IMAGE SIZE
        # =================================================

        file_size = os.path.getsize(
            image_path
        )

        print(
            "Image file size:",
            file_size,
            "bytes"
        )

        if file_size == 0:

            print(
                "ERROR: Image file is empty."
            )

            return None

        # =================================================
        # SUCCESS
        # =================================================

        print(
            "\n================================"
        )

        print(
            "IMAGE GENERATED SUCCESSFULLY"
        )

        print(
            "Saved:",
            os.path.abspath(
                image_path
            )
        )

        print(
            "Browser URL:"
        )

        print(
            f"/static/scenes/{image_filename}"
        )

        print(
            "Size:",
            file_size,
            "bytes"
        )

        print(
            "================================\n"
        )

        return image_path

    except Exception as e:

        print(
            "\n=============================="
        )

        print(
            "SCENE IMAGE ERROR"
        )

        print(
            "Error Type:",
            type(e).__name__
        )

        print(
            "Error:",
            str(e)
        )

        print(
            "=============================="
        )

        return None


# =====================================================
# HOME
# =====================================================

@app.route(
    "/",
    methods=["GET", "POST"]
)
def home():

    story = ""

    title = ""

    prompt = ""

    genre = "Fantasy"

    length = "Medium"

    tone = "Inspirational"

    character_name = ""

    number_of_characters = "1"

    setting = ""

    writing_style = "Creative"

    language = "English"

    custom_instructions = ""

    cover_path = None

    if request.method == "POST":

        prompt = request.form.get(
            "prompt",
            ""
        ).strip()

        genre = request.form.get(
            "genre",
            "Fantasy"
        )

        length = request.form.get(
            "length",
            "Medium"
        )

        tone = request.form.get(
            "tone",
            "Inspirational"
        )

        character_name = request.form.get(
            "character_name",
            ""
        ).strip()

        number_of_characters = request.form.get(
            "number_of_characters",
            "1"
        ).strip()

        setting = request.form.get(
            "setting",
            ""
        ).strip()

        writing_style = request.form.get(
            "writing_style",
            "Creative"
        ).strip()

        language = request.form.get(
            "language",
            "English"
        ).strip()

        custom_instructions = request.form.get(
            "custom_instructions",
            ""
        ).strip()

        title = generate_title(
            prompt,
            genre
        )

        story = generate_story(
            prompt,
            genre,
            length,
            tone,
            title,
            character_name,
            number_of_characters,
            setting,
            writing_style,
            language,
            custom_instructions
        )

        save_story(
            title,
            prompt,
            genre,
            length,
            tone,
            story,
            None
        )

        stories = get_stories()

        new_story_id = stories[0][0]

        cover_path = generate_cover(
            title,
            prompt,
            genre,
            tone,
            new_story_id
        )

        if cover_path:

            update_cover(
                new_story_id,
                cover_path
            )

    return render_template(
        "index.html",
        title=title,
        prompt=prompt,
        genre=genre,
        length=length,
        tone=tone,
        story=story,
        cover_path=cover_path,

        character_name=character_name,
        number_of_characters=number_of_characters,
        setting=setting,
        writing_style=writing_style,
        language=language,
        custom_instructions=custom_instructions
    )


# =====================================================
# REGENERATE STORY
# =====================================================

@app.route(
    "/regenerate",
    methods=["POST"]
)
def regenerate():

    prompt = request.form.get(
        "prompt",
        ""
    ).strip()

    genre = request.form.get(
        "genre",
        "Fantasy"
    )

    length = request.form.get(
        "length",
        "Medium"
    )

    tone = request.form.get(
        "tone",
        "Inspirational"
    )

    character_name = request.form.get(
        "character_name",
        ""
    ).strip()

    number_of_characters = request.form.get(
        "number_of_characters",
        "1"
    ).strip()

    setting = request.form.get(
        "setting",
        ""
    ).strip()

    writing_style = request.form.get(
        "writing_style",
        "Creative"
    ).strip()

    language = request.form.get(
        "language",
        "English"
    ).strip()

    custom_instructions = request.form.get(
        "custom_instructions",
        ""
    ).strip()

    title = generate_title(
        prompt,
        genre
    )

    story = generate_story(
        prompt,
        genre,
        length,
        tone,
        title,
        character_name,
        number_of_characters,
        setting,
        writing_style,
        language,
        custom_instructions
    )

    save_story(
        title,
        prompt,
        genre,
        length,
        tone,
        story,
        None
    )

    stories = get_stories()

    new_story_id = stories[0][0]

    cover_path = generate_cover(
        title,
        prompt,
        genre,
        tone,
        new_story_id
    )

    if cover_path:

        update_cover(
            new_story_id,
            cover_path
        )

    return render_template(
        "index.html",
        title=title,
        prompt=prompt,
        genre=genre,
        length=length,
        tone=tone,
        story=story,
        cover_path=cover_path,

        character_name=character_name,
        number_of_characters=number_of_characters,
        setting=setting,
        writing_style=writing_style,
        language=language,
        custom_instructions=custom_instructions
    )


# =====================================================
# STORY HISTORY
# =====================================================

@app.route("/history")
def history():

    stories = get_stories()

    audio_files = {}

    scene_files = {}

    for story in stories:

        audio_path = os.path.join(
            AUDIO_FOLDER,
            f"story_{story[0]}.wav"
        )

        audio_files[story[0]] = os.path.exists(
            audio_path
        )

        scene_path = os.path.join(
            SCENE_FOLDER,
            f"story_{story[0]}.json"
        )

        scene_files[story[0]] = os.path.exists(
            scene_path
        )

    return render_template(
        "history.html",
        stories=stories,
        audio_files=audio_files,
        scene_files=scene_files
    )


# =====================================================
# EDIT STORY
# =====================================================

@app.route(
    "/edit/<int:story_id>",
    methods=["GET"]
)
def edit_story(story_id):

    story_data = get_story(
        story_id
    )

    if not story_data:

        return redirect(
            url_for("history")
        )

    return render_template(
        "edit_story.html",
        story=story_data
    )


# =====================================================
# UPDATE EDITED STORY
# =====================================================

@app.route(
    "/edit/<int:story_id>",
    methods=["POST"]
)
def update_edited_story(story_id):

    title = request.form.get(
        "title",
        ""
    ).strip()

    story = request.form.get(
        "story",
        ""
    ).strip()

    update_story(
        story_id,
        title,
        story
    )

    return redirect(
        url_for("history")
    )


# =====================================================
# GENERATE COVER FOR OLD STORY
# =====================================================

@app.route(
    "/generate-cover/<int:story_id>",
    methods=["POST"]
)
def generate_old_cover(story_id):

    story_data = get_story(
        story_id
    )

    if not story_data:

        return redirect(
            url_for("history")
        )

    story_id = story_data[0]

    title = story_data[1]

    prompt = story_data[2]

    genre = story_data[3]

    tone = story_data[5]

    cover_path = generate_cover(
        title,
        prompt,
        genre,
        tone,
        story_id
    )

    if cover_path:

        update_cover(
            story_id,
            cover_path
        )

    return redirect(
        url_for("history")
    )


# =====================================================
# GENERATE STORY VOICE
# =====================================================

@app.route(
    "/generate_voice/<int:story_id>",
    methods=["POST"]
)
def generate_voice(story_id):

    print(
        "\n=============================="
    )

    print(
        "LOCAL VOICE GENERATION STARTED"
    )

    print(
        "Story ID:",
        story_id
    )

    print(
        "=============================="
    )

    try:

        story_data = get_story(
            story_id
        )

        if not story_data:

            print(
                "ERROR: Story not found"
            )

            return redirect(
                url_for("history")
            )

        story_text = story_data[6]

        if not story_text:

            print(
                "ERROR: Story text is empty"
            )

            return redirect(
                url_for("history")
            )

        audio_folder = os.path.join(
            "static",
            "audio"
        )

        os.makedirs(
            audio_folder,
            exist_ok=True
        )

        audio_filename = (
            f"story_{story_id}.wav"
        )

        audio_path = os.path.join(
            audio_folder,
            audio_filename
        )

        print(
            "Generating local voice..."
        )

        print(
            "Audio path:",
            os.path.abspath(
                audio_path
            )
        )

        engine = pyttsx3.init()

        engine.setProperty(
            "rate",
            160
        )

        engine.setProperty(
            "volume",
            1.0
        )

        engine.save_to_file(
            story_text,
            audio_path
        )

        engine.runAndWait()

        if os.path.exists(
            audio_path
        ):

            file_size = os.path.getsize(
                audio_path
            )

            print(
                "VOICE GENERATED SUCCESSFULLY!"
            )

            print(
                "File:",
                os.path.abspath(
                    audio_path
                )
            )

            print(
                "Size:",
                file_size,
                "bytes"
            )

        else:

            print(
                "ERROR: Audio file was not created"
            )

    except Exception as e:

        print(
            "=============================="
        )

        print(
            "LOCAL VOICE ERROR"
        )

        print(
            type(e).__name__
        )

        print(
            str(e)
        )

        print(
            "=============================="
        )

    return redirect(
        url_for("history")
    )


# =====================================================
# GENERATE AI SCENES FOR OLD STORY
# =====================================================

# =====================================================
# GENERATE AI SCENES FOR OLD STORY
# =====================================================

@app.route(
    "/generate-scenes/<int:story_id>",
    methods=["POST"]
)
def generate_scenes_route(story_id):

    print(
        "\n================================"
    )

    print(
        "GENERATE AI SCENES ROUTE"
    )

    print(
        "Story ID:",
        story_id
    )

    print(
        "================================"
    )

    # =================================================
    # GET STORY
    # =================================================

    story_data = get_story(
        story_id
    )

    if not story_data:

        print(
            "ERROR: Story not found."
        )

        return redirect(
            url_for("history")
        )

    # =================================================
    # GET STORY INFORMATION
    # =================================================

    title = story_data[1]
    story_text = story_data[6]
    genre = story_data[3]
    tone = story_data[5]

    if not story_text:

        print(
            "ERROR: Story text is empty."
        )

        return redirect(
            url_for("history")
        )

    # =================================================
    # GENERATE SCENES
    # =================================================

    try:

        generate_story_scenes(
            story_id,
            title,
            story_text,
            genre,
            tone
        )

        print(
            "AI scenes generated successfully."
        )

    except Exception as e:

        print(
            "\n=============================="
        )

        print(
            "SCENE GENERATION ERROR"
        )

        print(
            "Error Type:",
            type(e).__name__
        )

        print(
            "Error:",
            str(e)
        )

        print(
            "=============================="
        )

    # =================================================
    # SHOW SCENES PAGE
    # =================================================

    return redirect(
        url_for(
            "show_scenes",
            story_id=story_id
        )
    )


# =====================================================
# SHOW GENERATED SCENES
# =====================================================

# =====================================================
# SHOW GENERATED SCENES
# =====================================================

@app.route("/scenes/<int:story_id>", methods=["GET"])
def show_scenes(story_id):

    # =========================================
    # GET STORY
    # =========================================

    story_data = get_story(story_id)

    # Story nahi mili
    if not story_data:
        return redirect(url_for("history"))


    # =========================================
    # GET SCENES
    # =========================================

    scenes_data = get_story_scenes(story_id)


    # =========================================
    # NO SCENES
    # =========================================

    if not scenes_data:

        return render_template(
            "scenes.html",
            story=story_data,
            scenes=None,
            image_files={}
        )


    # =========================================
    # IMAGE FILE STATUS
    # =========================================

    image_files = {}


    # Get scenes list safely
    scenes_list = scenes_data.get(
        "scenes",
        []
    )


    # =========================================
    # CHECK EACH SCENE IMAGE
    # =========================================

    for index, scene in enumerate(
        scenes_list,
        start=1
    ):

        # Scene number
        scene_number = scene.get(
            "scene_number",
            index
        )


        # Convert scene number to integer
        try:

            scene_number = int(
                scene_number
            )

        except (TypeError, ValueError):

            scene_number = index


        # Image filename
        image_filename = (
            f"story_{story_id}"
            f"_scene_{scene_number}.png"
        )


        # Complete image path
        image_path = os.path.join(
            SCENE_FOLDER,
            image_filename
        )


        # Check image exists
        image_files[scene_number] = (
            os.path.exists(image_path)
        )


    # =========================================
    # SHOW SCENES PAGE
    # =========================================

    return render_template(
        "scenes.html",
        story=story_data,
        scenes=scenes_data,
        image_files=image_files
    )


# =====================================================
# GENERATE ONE SCENE IMAGE
# =====================================================

@app.route(
    "/generate-scene-image/<int:story_id>/<int:scene_number>",
    methods=["POST"]
)
def generate_scene_image_route(
    story_id,
    scene_number
):

    print(
        "\n================================"
    )

    print(
        "GENERATE SCENE IMAGE ROUTE"
    )

    print(
        "Story ID:",
        story_id
    )

    print(
        "Scene Number:",
        scene_number
    )

    print(
        "================================"
    )

    # =================================================
    # GET STORY
    # =================================================

    story_data = get_story(
        story_id
    )

    if not story_data:

        print(
            "ERROR: Story not found."
        )

        return redirect(
            url_for("history")
        )

    # =================================================
    # GET SCENES
    # =================================================

    scenes_data = get_story_scenes(
        story_id
    )

    if not scenes_data:

        print(
            "ERROR: No scenes found."
        )

        return redirect(
            url_for(
                "show_scenes",
                story_id=story_id
            )
        )

    # =================================================
    # GET SCENE LIST
    # =================================================

    scenes = scenes_data.get(
        "scenes",
        []
    )

    selected_scene = None

    # =================================================
    # FIND SELECTED SCENE
    # =================================================

    for scene in scenes:

        try:

            current_scene_number = int(
                scene.get(
                    "scene_number",
                    0
                )
            )

        except (
            ValueError,
            TypeError
        ):

            current_scene_number = 0

        if current_scene_number == scene_number:

            selected_scene = scene

            break

    # =================================================
    # SCENE NOT FOUND
    # =================================================

    if not selected_scene:

        print(
            "ERROR: Selected scene not found."
        )

        return redirect(
            url_for(
                "show_scenes",
                story_id=story_id
            )
        )

    # =================================================
    # GENERATE IMAGE
    # =================================================

    image_path = generate_scene_image(
        story_id,
        scene_number,
        selected_scene
    )

    # =================================================
    # RESULT
    # =================================================

    if image_path:

        print(
            "Scene image generated successfully."
        )

    else:

        print(
            "Scene image generation failed."
        )

    # =================================================
    # RETURN TO SCENES
    # =================================================

    return redirect(
        url_for(
            "show_scenes",
            story_id=story_id
        )
    )


# =====================================================
# DELETE STORY
# =====================================================

@app.route(
    "/delete/<int:story_id>",
    methods=["POST"]
)
def delete_story_route(story_id):

    story_data = get_story(
        story_id
    )

    # =================================================
    # DELETE COVER FILE
    # =================================================

    if story_data:

        cover_path = story_data[7]

        if cover_path:

            cover_file = os.path.join(
                "static",
                cover_path
            )

            if os.path.exists(
                cover_file
            ):

                try:

                    os.remove(
                        cover_file
                    )

                    print(
                        "Cover deleted:",
                        cover_file
                    )

                except Exception as e:

                    print(
                        "Cover delete error:",
                        e
                    )

    # =================================================
    # DELETE AUDIO FILES
    # =================================================

    audio_files_to_delete = [

        os.path.join(
            AUDIO_FOLDER,
            f"story_{story_id}.mp3"
        ),

        os.path.join(
            AUDIO_FOLDER,
            f"story_{story_id}.wav"
        )
    ]

    for audio_file in audio_files_to_delete:

        if os.path.exists(
            audio_file
        ):

            try:

                os.remove(
                    audio_file
                )

                print(
                    "Audio deleted:",
                    audio_file
                )

            except Exception as e:

                print(
                    "Audio delete error:",
                    e
                )

    # =================================================
    # DELETE SCENE JSON
    # =================================================

    scene_file = os.path.join(
        SCENE_FOLDER,
        f"story_{story_id}.json"
    )

    if os.path.exists(
        scene_file
    ):

        try:

            os.remove(
                scene_file
            )

            print(
                "Scene JSON deleted:",
                scene_file
            )

        except Exception as e:

            print(
                "Scene JSON delete error:",
                e
            )

    # =================================================
    # DELETE ALL SCENE IMAGES
    # =================================================

    if os.path.exists(
        SCENE_FOLDER
    ):

        scene_image_prefix = (
            f"story_{story_id}_scene_"
        )

        for filename in os.listdir(
            SCENE_FOLDER
        ):

            if (
                filename.startswith(
                    scene_image_prefix
                )
                and filename.lower().endswith(
                    ".png"
                )
            ):

                image_file = os.path.join(
                    SCENE_FOLDER,
                    filename
                )

                try:

                    os.remove(
                        image_file
                    )

                    print(
                        "Scene image deleted:",
                        image_file
                    )

                except Exception as e:

                    print(
                        "Scene image delete error:",
                        e
                    )

    # =================================================
    # DELETE DATABASE RECORD
    # =================================================

    delete_story(
        story_id
    )

    return redirect(
        url_for("history")
    )


# =====================================================
# DOWNLOAD TXT
# =====================================================

@app.route(
    "/download/txt",
    methods=["POST"]
)
def download_txt():

    story = request.form.get(
        "story",
        ""
    )

    title = request.form.get(
        "title",
        ""
    )

    prompt = request.form.get(
        "prompt",
        ""
    )

    genre = request.form.get(
        "genre",
        ""
    )

    length = request.form.get(
        "length",
        ""
    )

    tone = request.form.get(
        "tone",
        ""
    )

    content = f"""
AI STORY GENERATOR

Title:
{title}

Story Idea:
{prompt}

Genre:
{genre}

Length:
{length}

Tone:
{tone}

----------------------------------------

{story}
"""

    file = io.BytesIO()

    file.write(
        content.encode(
            "utf-8"
        )
    )

    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="ai_story.txt",
        mimetype="text/plain"
    )


# =====================================================
# DOWNLOAD WORD
# =====================================================

@app.route(
    "/download/word",
    methods=["POST"]
)
def download_word():

    story = request.form.get(
        "story",
        ""
    )

    title = request.form.get(
        "title",
        ""
    )

    prompt = request.form.get(
        "prompt",
        ""
    )

    genre = request.form.get(
        "genre",
        ""
    )

    length = request.form.get(
        "length",
        ""
    )

    tone = request.form.get(
        "tone",
        ""
    )

    document = Document()

    document.add_heading(
        "AI Story Generator",
        level=1
    )

    document.add_heading(
        title,
        level=2
    )

    document.add_paragraph(
        f"Story Idea: {prompt}"
    )

    document.add_paragraph(
        f"Genre: {genre}"
    )

    document.add_paragraph(
        f"Length: {length}"
    )

    document.add_paragraph(
        f"Tone: {tone}"
    )

    document.add_paragraph(
        "----------------------------------------"
    )

    document.add_heading(
        "Generated Story",
        level=2
    )

    document.add_paragraph(
        story
    )

    file = io.BytesIO()

    document.save(
        file
    )

    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="ai_story.docx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )


# =====================================================
# DOWNLOAD PDF
# =====================================================

@app.route(
    "/download/pdf",
    methods=["POST"]
)
def download_pdf():

    story = request.form.get(
        "story",
        ""
    )

    title = request.form.get(
        "title",
        ""
    )

    prompt = request.form.get(
        "prompt",
        ""
    )

    genre = request.form.get(
        "genre",
        ""
    )

    length = request.form.get(
        "length",
        ""
    )

    tone = request.form.get(
        "tone",
        ""
    )

    file = io.BytesIO()

    pdf = canvas.Canvas(
        file,
        pagesize=A4
    )

    width, height = A4

    y = height - 50

    pdf.setFont(
        "Helvetica-Bold",
        20
    )

    pdf.drawString(
        50,
        y,
        "AI Story Generator"
    )

    y -= 35

    pdf.setFont(
        "Helvetica-Bold",
        16
    )

    pdf.drawString(
        50,
        y,
        title[:80]
    )

    y -= 35

    pdf.setFont(
        "Helvetica",
        11
    )

    information = [

        f"Story Idea: {prompt}",

        f"Genre: {genre}",

        f"Length: {length}",

        f"Tone: {tone}"
    ]

    for info in information:

        pdf.drawString(
            50,
            y,
            info[:100]
        )

        y -= 20

    y -= 15

    pdf.setFont(
        "Helvetica-Bold",
        15
    )

    pdf.drawString(
        50,
        y,
        "Generated Story"
    )

    y -= 30

    pdf.setFont(
        "Helvetica",
        11
    )

    lines = story.split(
        "\n"
    )

    for paragraph in lines:

        words = paragraph.split()

        current_line = ""

        for word in words:

            test_line = (
                current_line +
                " " +
                word
            ).strip()

            if pdf.stringWidth(
                test_line,
                "Helvetica",
                11
            ) < 500:

                current_line = test_line

            else:

                pdf.drawString(
                    50,
                    y,
                    current_line
                )

                y -= 18

                current_line = word

                if y < 50:

                    pdf.showPage()

                    pdf.setFont(
                        "Helvetica",
                        11
                    )

                    y = height - 50

        if current_line:

            pdf.drawString(
                50,
                y,
                current_line
            )

            y -= 18

        y -= 8

        if y < 50:

            pdf.showPage()

            pdf.setFont(
                "Helvetica",
                11
            )

            y = height - 50

    pdf.save()

    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="ai_story.pdf",
        mimetype="application/pdf"
    )


# =====================================================
# RUN APPLICATION
# =====================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )
